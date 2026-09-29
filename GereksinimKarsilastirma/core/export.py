"""Excel çıktısı: DOORS'ta kullanılacak iş listesi ve raporlar."""
from __future__ import annotations

import io
from datetime import datetime

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

SEV_FILL = {
    "Yüksek": PatternFill("solid", fgColor="F8CBAD"),
    "Orta": PatternFill("solid", fgColor="FFE699"),
    "Düşük": PatternFill("solid", fgColor="DDEBF7"),
}
MATRIX_FILL = {
    "✓": PatternFill("solid", fgColor="C6EFCE"),
    "≠": PatternFill("solid", fgColor="F8CBAD"),
    "–": PatternFill("solid", fgColor="EDEDED"),
}


def diff_table(analysis: dict, project: dict) -> pd.DataFrame:
    dec, notes = project["decisions"], project["llm_notes"]
    rows = []
    for d in analysis["diffs"]:
        de = dec.get(d["key"], {})
        ln = notes.get(d["key"], {})
        rows.append({
            "Grup No": d.get("grup_no"),
            "Set": d["set"],
            "Req ID": d["id"],
            "Alan": d["alan"],
            "Kategori": d["kategori"],
            "Önem": d["onem"],
            "Detay": d["etiketler"],
            "Referans (çoğunluk/ana set)": d["referans"],
            "Bu setteki değer": d["deger"],
            "Referans belirsiz": "Evet" if d.get("belirsiz") else "",
            "Durum": de.get("durum", "Açık"),
            "Not": de.get("not", ""),
            "LLM açıklama": ln.get("aciklama", ""),
            "LLM önerisi": ln.get("oneri", ""),
            "key": d["key"],
        })
    cols = ["Grup No", "Set", "Req ID", "Alan", "Kategori", "Önem", "Detay",
            "Referans (çoğunluk/ana set)", "Bu setteki değer", "Referans belirsiz",
            "Durum", "Not", "LLM açıklama", "LLM önerisi", "key"]
    df = pd.DataFrame(rows, columns=cols)
    if not df.empty:
        order = {"Yüksek": 0, "Orta": 1, "Düşük": 2}
        df = df.sort_values(["Önem", "Grup No", "Set"], key=lambda s: s.map(order) if s.name == "Önem" else s,
                            na_position="last")
    return df


def matrix_table(analysis: dict) -> pd.DataFrame:
    sets = analysis["set_names"]
    real = {}
    for d in analysis["diffs"]:
        if d.get("grup") and not d["alan"].startswith("Kontrol:"):
            real.setdefault(d["grup"], set()).add(d["set"])
    rows = []
    for c in analysis["clusters"]:
        if not c["common"]:
            continue
        r = {"Grup No": c["no"], "Başlık / Metin": c["title"], "Kapsam": f"{c['coverage']}/{len(sets)}",
             "En yüksek önem": c["max_sev"]}
        bad = real.get(c["cid"], set())
        for s in sets:
            r[s] = "–" if s not in c["sets"] else ("≠" if s in bad else "✓")
        rows.append(r)
    return pd.DataFrame(rows)


def common_table(analysis: dict) -> pd.DataFrame:
    items = analysis["items"]
    bad_items = {d["item"] for d in analysis["diffs"] if d.get("item") and not d["alan"].startswith("Kontrol:")}
    rows = []
    for c in analysis["clusters"]:
        if not c["common"]:
            continue
        for i in c["members"]:
            it = items[i]
            r = {"Grup No": c["no"], "Set": it["set"], "Ekipman": it["equipment"], "Req ID": it["id"],
                 "Source": it["source"], "Metin": it["text"], "Attribute": it["type"]}
            r.update(it["attrs"])
            r["Bölüm"] = it["section"]
            r["Fark var"] = "Evet" if it["key"] in bad_items else ""
            rows.append(r)
    return pd.DataFrame(rows)


def single_table(analysis: dict) -> pd.DataFrame:
    items = analysis["items"]
    rows = []
    for c in analysis["clusters"]:
        if c["common"]:
            continue
        for i in c["members"]:
            it = items[i]
            rows.append({"Set": it["set"], "Req ID": it["id"], "Source": it["source"],
                         "Attribute": it["type"], "Metin": it["text"]})
    return pd.DataFrame(rows)


def _format(ws, df: pd.DataFrame, wide_cols=()):
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for idx, col in enumerate(df.columns, start=1):
        letter = ws.cell(row=1, column=idx).column_letter
        width = 60 if col in wide_cols else min(max(len(str(col)) + 2, 12), 30)
        ws.column_dimensions[letter].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def to_excel(analysis: dict, project: dict, verify_rows: list[dict] | None = None) -> bytes:
    buf = io.BytesIO()
    diffs = diff_table(analysis, project).drop(columns=["key"])
    matrix = matrix_table(analysis)
    common = common_table(analysis)
    single = single_table(analysis)
    n_common = sum(1 for c in analysis["clusters"] if c["common"])
    summary = pd.DataFrame([
        ("Proje", project["name"]),
        ("Rapor tarihi", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Set sayısı", len(analysis["set_names"])),
        ("Ortak gereksinim grubu", n_common),
        ("Tek sette kalan öğe", len(single)),
        ("Toplam fark / bulgu", len(diffs)),
        ("Yüksek önem", int((diffs["Önem"] == "Yüksek").sum()) if not diffs.empty else 0),
        ("Orta önem", int((diffs["Önem"] == "Orta").sum()) if not diffs.empty else 0),
        ("Düşük önem", int((diffs["Önem"] == "Düşük").sum()) if not diffs.empty else 0),
    ], columns=["Bilgi", "Değer"])

    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="Özet", index=False)
        diffs.to_excel(xw, sheet_name="Fark Listesi", index=False)
        matrix.to_excel(xw, sheet_name="Kapsam Matrisi", index=False)
        common.to_excel(xw, sheet_name="Ortak Gereksinimler", index=False)
        single.to_excel(xw, sheet_name="Tek Sette Olanlar", index=False)
        if verify_rows:
            vdf = pd.DataFrame(verify_rows).drop(columns=["key"])
            vdf.to_excel(xw, sheet_name="Doğrulama", index=False)
        wb = xw.book
        _format(wb["Özet"], summary, wide_cols=("Değer",))
        wb["Özet"].column_dimensions["A"].width = 28
        _format(wb["Fark Listesi"], diffs, wide_cols=("Referans (çoğunluk/ana set)", "Bu setteki değer",
                                                      "LLM açıklama", "LLM önerisi", "Detay"))
        ws = wb["Fark Listesi"]
        sev_col = list(diffs.columns).index("Önem") + 1
        for r in range(2, ws.max_row + 1):
            c = ws.cell(row=r, column=sev_col)
            if c.value in SEV_FILL:
                c.fill = SEV_FILL[c.value]
        _format(wb["Kapsam Matrisi"], matrix, wide_cols=("Başlık / Metin",))
        ws = wb["Kapsam Matrisi"]
        for row in ws.iter_rows(min_row=2, min_col=5):
            for c in row:
                if c.value in MATRIX_FILL:
                    c.fill = MATRIX_FILL[c.value]
                    c.alignment = Alignment(horizontal="center")
        _format(wb["Ortak Gereksinimler"], common, wide_cols=("Metin", "Rationale", "Notes"))
        _format(wb["Tek Sette Olanlar"], single, wide_cols=("Metin",))
        if verify_rows:
            _format(wb["Doğrulama"], vdf, wide_cols=("Önceki değer", "Güncel değer", "Referans"))
    return buf.getvalue()

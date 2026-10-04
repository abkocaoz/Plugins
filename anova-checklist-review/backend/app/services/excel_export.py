"""Phase 5: fill a *copy* of the checklist Excel template (never mutate the original)."""

from __future__ import annotations

import hashlib
import shutil
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import Settings
from app.models.entities import (
    ChecklistAnswer,
    ChecklistDefinition,
    ChecklistItem,
    ChecklistRun,
    Export,
    ExtractedReference,
    ReferenceFinding,
    ReferenceMatch,
    StandardVersion,
)
from app.services.checklist_catalog import CATALOG_DIR

# States that must leave Yes/No/NA answer cells blank (explanation goes to Comment).
BLANK_ANSWER_STATES = frozenset({"INSUFFICIENT_EVIDENCE", "MANUAL_REVIEW", "ERROR", "PENDING"})

ANSWER_DISPLAY = {"YES": "Yes", "NO": "No", "NA": "NA"}

REF_SHEET_NAME = "Reference Validation"
REF_HEADERS = (
    "ref",
    "stated version",
    "selected source",
    "result",
    "location",
    "explanation",
)


def sanitize_excel_text(value: str | None) -> str | None:
    """Guard formula injection: prefix external text that Excel may treat as formula."""
    if value is None:
        return None
    text = str(value)
    if not text:
        return text
    if text[0] in ("=", "+", "-", "@"):
        return "'" + text
    return text


def resolve_template_path(definition: ChecklistDefinition) -> Path:
    raw = definition.excel_template_path
    if not raw:
        raise FileNotFoundError(
            "checklist definition has no excel_template_path; supply a production template "
            "or use the synthetic fixture path in the catalog"
        )
    candidates = [
        Path(raw),
        CATALOG_DIR / raw,
        CATALOG_DIR / Path(raw).name,
        Path(__file__).resolve().parents[2] / raw,
    ]
    for cand in candidates:
        if cand.is_file():
            return cand.resolve()
    raise FileNotFoundError(f"Excel template not found for path={raw!r}; tried {[str(c) for c in candidates]}")


def answer_cell_value(state: str | None, *, answer_values: dict[str, str] | None = None) -> str | None:
    """Map engine state → template Yes/No/NA; blank for insufficient/manual/error."""
    if not state:
        return None
    blank = BLANK_ANSWER_STATES
    mapping_meta = answer_values or ANSWER_DISPLAY
    if state in blank:
        return None
    if state in mapping_meta:
        return mapping_meta[state]
    # Unknown non-YN states stay blank rather than inventing values
    if state not in {"YES", "NO", "NA"}:
        return None
    return ANSWER_DISPLAY.get(state)


@dataclass
class ExportRowWrite:
    item_key: str
    cells: dict[str, str | None]
    state_used: str | None


def _effective_state(answer: ChecklistAnswer, mode: str) -> str | None:
    if mode == "reviewer_approved":
        return answer.human_decision or None
    # draft: prefer human if already set, else AI proposal / effective state
    if answer.human_decision:
        return answer.human_decision
    prop = (answer.ai_proposal or {}).get("proposed_state")
    return prop or answer.state


def _comment_for_export(
    answer: ChecklistAnswer,
    state: str | None,
    ref_summary: str | None,
) -> str:
    parts: list[str] = []
    if answer.comment_text:
        parts.append(answer.comment_text)
    elif answer.rationale:
        parts.append(answer.rationale)
    if state in BLANK_ANSWER_STATES:
        parts.append(f"[Answer left blank: {state}]")
    if ref_summary:
        parts.append(f"Related refs: {ref_summary}")
    return " | ".join(p for p in parts if p)


async def _reference_rows_for_document(
    session: AsyncSession, document_version_id: uuid.UUID
) -> list[dict[str, str]]:
    refs = (
        await session.scalars(
            select(ExtractedReference).where(
                ExtractedReference.document_version_id == document_version_id
            )
        )
    ).all()
    rows: list[dict[str, str]] = []
    for ref in refs:
        findings = (
            await session.scalars(
                select(ReferenceFinding).where(ReferenceFinding.extracted_reference_id == ref.id)
            )
        ).all()
        selected_source = ""
        if ref.selected_match_id:
            match = await session.get(ReferenceMatch, ref.selected_match_id)
            if match and match.standard_version_id:
                sv = await session.get(StandardVersion, match.standard_version_id)
                if sv:
                    selected_source = f"{sv.version_label or ''} ({sv.id})".strip()
                else:
                    selected_source = str(match.standard_version_id)
            elif match:
                selected_source = (match.meta or {}).get("label") or str(match.id)
        loc = ref.locator or {}
        location = (
            loc.get("path")
            or loc.get("page")
            or loc.get("sheet")
            or loc.get("unit_id")
            or json_compact(loc)
        )
        if findings:
            for f in findings:
                # Short explanation only — do not republish long standard text
                explanation = (f.message or "")[:500]
                rows.append(
                    {
                        "ref": (ref.doc_id_guess or ref.raw_text or "")[:300],
                        "stated version": (ref.version_guess or "")[:120],
                        "selected source": selected_source[:300],
                        "result": f.status,
                        "location": str(location)[:300],
                        "explanation": explanation,
                    }
                )
        else:
            rows.append(
                {
                    "ref": (ref.doc_id_guess or ref.raw_text or "")[:300],
                    "stated version": (ref.version_guess or "")[:120],
                    "selected source": selected_source[:300],
                    "result": ref.resolution_state or "",
                    "location": str(location)[:300],
                    "explanation": (ref.context_span or "")[:200],
                }
            )
    return rows


def json_compact(obj: Any) -> str:
    if not obj:
        return ""
    if isinstance(obj, dict):
        return ",".join(f"{k}={v}" for k, v in list(obj.items())[:6])
    return str(obj)[:200]


def write_reference_validation_sheet(wb, rows: list[dict[str, str]]) -> None:
    if REF_SHEET_NAME in wb.sheetnames:
        ws = wb[REF_SHEET_NAME]
        # Clear prior data rows but keep sheet object
        if ws.max_row and ws.max_row > 0:
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=len(REF_HEADERS)):
                for cell in row:
                    cell.value = None
    else:
        ws = wb.create_sheet(REF_SHEET_NAME)
    for col, header in enumerate(REF_HEADERS, 1):
        ws.cell(1, col, header)
    for r_idx, row in enumerate(rows, 2):
        for c_idx, key in enumerate(REF_HEADERS, 1):
            ws.cell(r_idx, c_idx, sanitize_excel_text(row.get(key, "")))


def fill_workbook_copy(
    template_path: Path,
    dest_path: Path,
    *,
    item_writes: list[ExportRowWrite],
    reference_rows: list[dict[str, str]],
    cell_mapping: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Copy template → dest, fill mapped cells, add Reference Validation sheet."""
    dest_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, dest_path)

    wb = load_workbook(dest_path)
    sheet_name = (cell_mapping or {}).get("sheet") or "Checklist"
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Template missing sheet {sheet_name!r}")
    ws = wb[sheet_name]

    # Snapshot preservation probes before write
    merged_before = list(ws.merged_cells.ranges)
    validations_before = len(ws.data_validations.dataValidation) if ws.data_validations else 0
    print_area_before = ws.print_area
    formula_j3_before = ws["J3"].value

    written = 0
    for row_write in item_writes:
        for addr, value in row_write.cells.items():
            if not addr:
                continue
            ws[addr] = sanitize_excel_text(value) if isinstance(value, str) else value
            written += 1

    write_reference_validation_sheet(wb, reference_rows)
    wb.save(dest_path)
    wb.close()

    # Re-open to confirm preservation
    check = load_workbook(dest_path)
    ws2 = check[sheet_name]
    merged_after = list(ws2.merged_cells.ranges)
    validations_after = len(ws2.data_validations.dataValidation) if ws2.data_validations else 0
    print_area_after = ws2.print_area
    formula_j3_after = ws2["J3"].value
    check.close()

    digest = hashlib.sha256(dest_path.read_bytes()).hexdigest()
    return {
        "cells_written": written,
        "reference_rows": len(reference_rows),
        "content_sha256": digest,
        "preservation": {
            "merged_cells_ok": [str(m) for m in merged_before] == [str(m) for m in merged_after],
            "data_validations_ok": validations_before == validations_after,
            "print_area_ok": print_area_before == print_area_after,
            "formula_j3_ok": formula_j3_before == formula_j3_after,
        },
        "template_path": str(template_path),
        "dest_path": str(dest_path),
    }


async def build_export_item_writes(
    session: AsyncSession,
    run: ChecklistRun,
    definition: ChecklistDefinition,
    mode: str,
) -> tuple[list[ExportRowWrite], list[dict[str, str]]]:
    cell_mapping = (definition.meta or {}).get("cell_mapping") or {}
    answer_values = cell_mapping.get("answer_values") or ANSWER_DISPLAY
    blank_states = set(cell_mapping.get("blank_answer_states") or BLANK_ANSWER_STATES)

    items = {
        i.id: i
        for i in (
            await session.scalars(
                select(ChecklistItem).where(ChecklistItem.definition_id == definition.id)
            )
        ).all()
    }
    answers = (
        await session.scalars(select(ChecklistAnswer).where(ChecklistAnswer.run_id == run.id))
    ).all()

    # Compact related reference findings for comments (not full standard text)
    ref_rows = await _reference_rows_for_document(session, run.document_version_id)
    ref_summaries: list[str] = []
    for r in ref_rows[:12]:
        bit = f"{r.get('ref','')}:{r.get('result','')}"
        if bit.strip(":"):
            ref_summaries.append(bit[:80])
    ref_summary = "; ".join(ref_summaries) if ref_summaries else None

    writes: list[ExportRowWrite] = []
    for ans in answers:
        item = items.get(ans.item_id)
        if item is None:
            continue
        schema = item.schema_json or {}
        cells_map = schema.get("excel_cells") or {}
        state = _effective_state(ans, mode)
        human_export = (ans.raw_model_output or {}).get("human_export") or {}
        # Status / Reviewed Item: only human-supplied values — never Closed because AI answered
        status_val = human_export.get("status") if ans.human_decision else None
        reviewed_val = human_export.get("reviewed_item") if ans.human_decision else None

        if mode == "reviewer_approved" and not ans.human_decision:
            answer_val = None
            status_val = None
            reviewed_val = None
        elif state in blank_states:
            answer_val = None
        else:
            answer_val = answer_cell_value(state, answer_values=answer_values)

        comment = _comment_for_export(ans, state, ref_summary)
        chapter = ans.chapter_text
        references_cell = ref_summary
        # Applicability column is NEVER filled from conformity Answer
        is_applicable_val = (ans.ai_proposal or {}).get("is_applicable")
        if isinstance(is_applicable_val, str):
            is_applicable_val = is_applicable_val.strip() or None
        else:
            is_applicable_val = None

        cell_values: dict[str, str | None] = {}
        if cells_map.get("is_applicable"):
            cell_values[cells_map["is_applicable"]] = is_applicable_val
        if cells_map.get("answer"):
            cell_values[cells_map["answer"]] = answer_val
        if cells_map.get("chapter"):
            cell_values[cells_map["chapter"]] = chapter
        if cells_map.get("comment"):
            cell_values[cells_map["comment"]] = comment
        if cells_map.get("references"):
            cell_values[cells_map["references"]] = references_cell
        if cells_map.get("reviewed_item"):
            cell_values[cells_map["reviewed_item"]] = reviewed_val
        if cells_map.get("status"):
            cell_values[cells_map["status"]] = status_val

        writes.append(
            ExportRowWrite(
                item_key=item.item_key,
                cells=cell_values,
                state_used=state,
            )
        )
    return writes, ref_rows


async def run_export_job(
    session: AsyncSession, settings: Settings, job_payload: dict[str, Any]
) -> dict[str, Any]:
    export_id = uuid.UUID(job_payload["export_id"])
    export = await session.get(Export, export_id)
    if export is None:
        raise RuntimeError(f"export {export_id} not found")
    run = await session.get(ChecklistRun, export.checklist_run_id)
    if run is None:
        raise RuntimeError("checklist run missing for export")
    definition = await session.get(ChecklistDefinition, run.definition_id)
    if definition is None:
        raise RuntimeError("checklist definition missing")

    mode = (export.meta or {}).get("mode") or job_payload.get("mode") or "draft"
    template_path = resolve_template_path(definition)
    dest = Path(export.storage_path)
    if not dest.is_absolute():
        dest = Path(settings.export_dir) / dest

    item_writes, ref_rows = await build_export_item_writes(session, run, definition, mode)
    meta = fill_workbook_copy(
        template_path,
        dest,
        item_writes=item_writes,
        reference_rows=ref_rows,
        cell_mapping=(definition.meta or {}).get("cell_mapping"),
    )

    # Ensure original template untouched
    if template_path.resolve() == dest.resolve():
        raise RuntimeError("refusing to write export onto template path")

    export.status = "ready"
    export.content_sha256 = meta["content_sha256"]
    export.template_path = str(template_path)
    export.meta = {
        **(export.meta or {}),
        "mode": mode,
        "cells_written": meta["cells_written"],
        "reference_rows": meta["reference_rows"],
        "preservation": meta["preservation"],
        "item_keys": [w.item_key for w in item_writes],
    }
    await session.flush()
    return {"export_id": str(export.id), "status": export.status, **meta}

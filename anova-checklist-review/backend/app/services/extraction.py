"""Document text extraction scaffolding with source locators.

Supports PDF / DOCX / XLSX / CSV / plain code-ish text.
Does not perform OCR in Phase 2; records OCR_NEEDED / UNREADABLE issues instead.
"""

from __future__ import annotations

import csv
import io
import json
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

import chardet
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

CODE_SUFFIXES = {
    ".c",
    ".h",
    ".cpp",
    ".hpp",
    ".cc",
    ".py",
    ".js",
    ".ts",
    ".java",
    ".go",
    ".rs",
    ".cs",
    ".sql",
    ".sh",
    ".bash",
    ".yml",
    ".yaml",
    ".json",
    ".toml",
    ".md",
    ".txt",
    ".xml",
}


@dataclass
class Locator:
    kind: str
    page: int | None = None
    paragraph: int | None = None
    sheet: str | None = None
    row: int | None = None
    column: str | None = None
    line_start: int | None = None
    line_end: int | None = None
    extra: dict[str, Any] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        return {k: v for k, v in asdict(self).items() if v is not None and v != {}}


@dataclass
class ExtractedUnit:
    text: str
    locator: Locator


@dataclass
class ExtractionIssue:
    code: str
    message: str
    locator: dict[str, Any] = field(default_factory=dict)


@dataclass
class ExtractionResult:
    units: list[ExtractedUnit]
    issues: list[ExtractionIssue]
    format: str
    full_text: str

    @property
    def status(self) -> str:
        if not self.units and self.issues:
            return "failed"
        if self.issues:
            return "partial"
        return "ok"

    def to_json_dict(self) -> dict[str, Any]:
        return {
            "format": self.format,
            "status": self.status,
            "issues": [asdict(i) for i in self.issues],
            "units": [{"text": u.text, "locator": u.locator.to_dict()} for u in self.units],
            "full_text": self.full_text,
        }


def detect_format(filename: str, mime_type: str | None = None) -> str:
    suffix = Path(filename).suffix.lower()
    mime = (mime_type or "").lower()
    if suffix == ".pdf" or "pdf" in mime:
        return "pdf"
    if suffix in {".docx"} or "wordprocessingml" in mime:
        return "docx"
    if suffix in {".xlsx", ".xlsm"} or "spreadsheetml" in mime:
        return "xlsx"
    if suffix == ".csv" or "csv" in mime:
        return "csv"
    if suffix in CODE_SUFFIXES or mime.startswith("text/"):
        return "code"
    return "unknown"


def _decode_bytes(data: bytes) -> tuple[str, str]:
    guess = chardet.detect(data) or {}
    encoding = guess.get("encoding") or "utf-8"
    try:
        return data.decode(encoding), encoding
    except UnicodeDecodeError:
        return data.decode("utf-8", errors="replace"), "utf-8-replace"


def extract_pdf(data: bytes) -> ExtractionResult:
    from pypdf import PdfReader

    issues: list[ExtractionIssue] = []
    units: list[ExtractedUnit] = []
    reader = PdfReader(io.BytesIO(data))
    for idx, page in enumerate(reader.pages, start=1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception as exc:  # noqa: BLE001
            issues.append(
                ExtractionIssue(
                    code="UNREADABLE",
                    message=f"PDF page extract failed: {exc}",
                    locator={"kind": "pdf_page", "page": idx},
                )
            )
            continue
        if not text:
            issues.append(
                ExtractionIssue(
                    code="OCR_NEEDED",
                    message="No extractable text on page (likely scanned/image PDF)",
                    locator={"kind": "pdf_page", "page": idx},
                )
            )
            continue
        units.append(ExtractedUnit(text=text, locator=Locator(kind="pdf_page", page=idx)))
    full = "\n\n".join(u.text for u in units)
    return ExtractionResult(units=units, issues=issues, format="pdf", full_text=full)


def extract_docx(data: bytes) -> ExtractionResult:
    from docx import Document as DocxDocument

    issues: list[ExtractionIssue] = []
    units: list[ExtractedUnit] = []
    doc = DocxDocument(io.BytesIO(data))
    for idx, para in enumerate(doc.paragraphs, start=1):
        text = (para.text or "").strip()
        if text:
            units.append(
                ExtractedUnit(text=text, locator=Locator(kind="docx_paragraph", paragraph=idx))
            )
    # Tables as additional units
    for t_idx, table in enumerate(doc.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [((c.text or "").strip()) for c in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            units.append(
                ExtractedUnit(
                    text="\n".join(rows),
                    locator=Locator(kind="docx_table", extra={"table": t_idx}),
                )
            )
    if not units:
        issues.append(
            ExtractionIssue(
                code="UNREADABLE",
                message="DOCX contained no paragraph/table text",
                locator={"kind": "docx"},
            )
        )
    full = "\n".join(u.text for u in units)
    return ExtractionResult(units=units, issues=issues, format="docx", full_text=full)


def extract_xlsx(data: bytes) -> ExtractionResult:
    issues: list[ExtractionIssue] = []
    units: list[ExtractedUnit] = []
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        return ExtractionResult(
            units=[],
            issues=[ExtractionIssue(code="UNREADABLE", message=f"XLSX open failed: {exc}")],
            format="xlsx",
            full_text="",
        )
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=False):
            for cell in row:
                if cell.value is None:
                    continue
                text = str(cell.value).strip()
                if not text:
                    continue
                col = getattr(cell, "column_letter", None) or get_column_letter(cell.column)
                units.append(
                    ExtractedUnit(
                        text=text,
                        locator=Locator(
                            kind="xlsx_cell",
                            sheet=sheet.title,
                            row=cell.row,
                            column=str(col),
                        ),
                    )
                )
    wb.close()
    if not units:
        issues.append(
            ExtractionIssue(code="UNREADABLE", message="XLSX had no non-empty cells", locator={})
        )
    full = "\n".join(u.text for u in units)
    return ExtractionResult(units=units, issues=issues, format="xlsx", full_text=full)


def extract_csv(data: bytes) -> ExtractionResult:
    text, encoding = _decode_bytes(data)
    issues: list[ExtractionIssue] = []
    units: list[ExtractedUnit] = []
    try:
        sample = text[:4096]
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
        issues.append(
            ExtractionIssue(
                code="CSV_DIALECT_FALLBACK",
                message="Could not sniff CSV dialect; using excel defaults",
            )
        )
    reader = csv.reader(io.StringIO(text), dialect)
    for row_idx, row in enumerate(reader, start=1):
        line = ",".join(row).strip()
        if line:
            units.append(
                ExtractedUnit(
                    text=line,
                    locator=Locator(kind="csv_row", row=row_idx, extra={"encoding": encoding}),
                )
            )
    full = "\n".join(u.text for u in units)
    return ExtractionResult(units=units, issues=issues, format="csv", full_text=full)


def extract_code(data: bytes, filename: str) -> ExtractionResult:
    text, encoding = _decode_bytes(data)
    issues: list[ExtractionIssue] = []
    if "\x00" in text:
        issues.append(
            ExtractionIssue(
                code="UNREADABLE",
                message="Binary/null bytes detected; treating as unreadable code upload",
            )
        )
        return ExtractionResult(units=[], issues=issues, format="code", full_text="")
    units: list[ExtractedUnit] = []
    lines = text.splitlines()
    # Chunk by ~80 lines while preserving line locators
    block = 80
    for start in range(0, len(lines), block):
        chunk_lines = lines[start : start + block]
        body = "\n".join(chunk_lines).strip()
        if not body:
            continue
        units.append(
            ExtractedUnit(
                text=body,
                locator=Locator(
                    kind="code_lines",
                    line_start=start + 1,
                    line_end=start + len(chunk_lines),
                    extra={"filename": Path(filename).name, "encoding": encoding},
                ),
            )
        )
    if not units:
        issues.append(ExtractionIssue(code="UNREADABLE", message="Empty code/text file"))
    return ExtractionResult(
        units=units, issues=issues, format="code", full_text=text
    )


def extract_document(
    data: bytes, filename: str, mime_type: str | None = None
) -> ExtractionResult:
    fmt = detect_format(filename, mime_type)
    if fmt == "pdf":
        return extract_pdf(data)
    if fmt == "docx":
        return extract_docx(data)
    if fmt == "xlsx":
        return extract_xlsx(data)
    if fmt == "csv":
        return extract_csv(data)
    if fmt == "code":
        return extract_code(data, filename)
    return ExtractionResult(
        units=[],
        issues=[
            ExtractionIssue(
                code="UNSUPPORTED_FORMAT",
                message=f"Unsupported format for '{filename}' (mime={mime_type})",
            )
        ],
        format="unknown",
        full_text="",
    )


def dumps_extraction(result: ExtractionResult) -> bytes:
    return json.dumps(result.to_json_dict(), ensure_ascii=False, indent=2).encode("utf-8")

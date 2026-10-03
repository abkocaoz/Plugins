#!/usr/bin/env python3
"""Regenerate minimal synthetic DataICD .xlsx fixture (applicability ≠ conformity)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

OUT = (
    Path(__file__).resolve().parents[1]
    / "app"
    / "catalogs"
    / "fixtures"
    / "data_icd_synthetic_v1.xlsx"
)


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Checklist"

    ws.merge_cells("A1:I1")
    ws["A1"] = "DataICD Checklist (SYNTHETIC FIXTURE — not a production template)"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")

    # Is Applicable is separate from Answer (conformity) — never merge those roles
    headers = [
        "Question Number",
        "Question",
        "Is Applicable",
        "Answer",
        "Chapter",
        "Comment",
        "References",
        "Reviewed Item",
        "Status",
    ]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(2, col, h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E2EFDA")

    questions = [
        ("1", "Is a Data ICD document applicable to this project baseline?"),
        ("2", "Does every ICD record include a signal/parameter identifier?"),
        ("3", "Are signal/parameter identifiers unique across all ICD records?"),
        ("4", "Does every ICD record specify a data type?"),
        ("5", "Does every ICD record specify units (or N/A for dimensionless)?"),
        ("6", "Are ICD record descriptions present for all records (manual judgment)?"),
    ]
    thin = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    for i, (num, q) in enumerate(questions):
        row = 3 + i
        ws.cell(row, 1, num).border = thin
        ws.cell(row, 2, q).border = thin
        for c in range(3, 10):
            ws.cell(row, c, None).border = thin

    dv_app = DataValidation(type="list", formula1='"Yes,No,NA"', allow_blank=True)
    ws.add_data_validation(dv_app)
    dv_app.add("C3:C8")
    dv_ans = DataValidation(type="list", formula1='"Yes,No,NA"', allow_blank=True)
    ws.add_data_validation(dv_ans)
    dv_ans.add("D3:D8")
    dv_st = DataValidation(type="list", formula1='"Open,In Review,Closed"', allow_blank=True)
    ws.add_data_validation(dv_st)
    dv_st.add("I3:I8")

    ws["K2"] = "FilledAnswers"
    ws["K3"] = "=COUNTA(D3:D8)"
    wb.defined_names.add(DefinedName(name="DataIcdAnswerColumn", attr_text="Checklist!$D$3:$D$8"))
    ws.print_area = "A1:I8"
    ws.page_setup.orientation = "landscape"
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 70
    for col in "CDEFGHI":
        ws.column_dimensions[col].width = 16

    notes = wb.create_sheet("Notes")
    notes["A1"] = "Synthetic DataICD fixture. Is Applicable (col C) ≠ Answer/conformity (col D)."
    notes.merge_cells("A1:E1")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()

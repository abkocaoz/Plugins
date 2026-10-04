#!/usr/bin/env python3
"""Regenerate the minimal synthetic Software Code Standard .xlsx fixture."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

OUT = (
    Path(__file__).resolve().parents[1]
    / "app"
    / "catalogs"
    / "fixtures"
    / "software_code_standard_synthetic_v1.xlsx"
)


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Checklist"

    ws.merge_cells("A1:H1")
    ws["A1"] = (
        "Software Code Standard Checklist "
        "(SYNTHETIC FIXTURE — not a production template)"
    )
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")

    headers = [
        "Question Number",
        "Question",
        "Answer",
        "Chapter",
        "Comment",
        "References",
        "Reviewed Item",
        "Status",
    ]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(2, col, h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E1F2")

    questions = [
        ("1", "Does every source file include a file header comment with copyright or ownership notice?"),
        ("2", "Are TODO/FIXME markers absent from all production source pages/units?"),
        ("3", "Is the coding standard document itself applicable and cited for this software baseline?"),
        ("4", "Does the design/code documentation adequately describe naming conventions used in the software?"),
        ("5", "Are related DO-178C (or project coding) standard clauses available and consistent for this review?"),
        ("6", "Is there traceability from checklist-required coding rules to sections in the uploaded source document?"),
        ("7", "Has external static-analysis / peer-review evidence been provided for coding-standard compliance?"),
        ("8", "Are project-specific deviations from the coding standard documented and approved?"),
    ]
    thin = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    for i, (num, q) in enumerate(questions):
        row = 3 + i
        ws.cell(row, 1, num).border = thin
        ws.cell(row, 2, q).border = thin
        for c in range(3, 9):
            ws.cell(row, c, None).border = thin

    dv = DataValidation(type="list", formula1='"Yes,No,NA"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("C3:C10")
    dv2 = DataValidation(type="list", formula1='"Open,In Review,Closed"', allow_blank=True)
    ws.add_data_validation(dv2)
    dv2.add("H3:H10")

    ws["J2"] = "FilledAnswers"
    ws["J3"] = "=COUNTA(C3:C10)"
    wb.defined_names.add(DefinedName(name="AnswerColumn", attr_text="Checklist!$C$3:$C$10"))
    ws.print_area = "A1:H10"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 70
    for col in "CDEFGH":
        ws.column_dimensions[col].width = 18

    ws2 = wb.create_sheet("Notes")
    ws2["A1"] = "Synthetic fixture for Phase 5 mapping/preservation tests."
    ws2["A2"] = "Production templates must be supplied separately."
    ws2.merge_cells("A1:D1")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()

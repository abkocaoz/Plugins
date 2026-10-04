#!/usr/bin/env python3
"""Regenerate minimal synthetic SECI .xlsx fixture for cross-document checks."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

OUT = (
    Path(__file__).resolve().parents[1]
    / "app"
    / "catalogs"
    / "fixtures"
    / "seci_synthetic_v1.xlsx"
)


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Checklist"

    ws.merge_cells("A1:H1")
    ws["A1"] = "SECI Cross-Document Checklist (SYNTHETIC FIXTURE — not a production template)"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")

    headers = [
        "Question Number",
        "Question",
        "Answer",
        "Chapter",
        "Comment",
        "References",
        "Reviewed Item",
        "Status",
    ]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(2, col, h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FCE4D6")

    questions = [
        ("1", "Are SECI signal IDs present in the pinned Data ICD document?"),
        ("2", "Do all SECI shall-style requirements map to at least one peer document unit?"),
        ("3", "Are all pinned document versions available for cross-document comparison?"),
        ("4", "Is naming of shared identifiers consistent across SECI and design documents?"),
    ]
    thin = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    for i, (num, q) in enumerate(questions):
        row = 3 + i
        ws.cell(row, 1, num).border = thin
        ws.cell(row, 2, q).border = thin
        for c in range(3, 9):
            ws.cell(row, c, None).border = thin

    dv = DataValidation(type="list", formula1='"Yes,No,NA"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("C3:C6")
    dv2 = DataValidation(type="list", formula1='"Open,In Review,Closed"', allow_blank=True)
    ws.add_data_validation(dv2)
    dv2.add("H3:H6")

    ws["J2"] = "FilledAnswers"
    ws["J3"] = "=COUNTA(C3:C6)"
    wb.defined_names.add(DefinedName(name="SeciAnswerColumn", attr_text="Checklist!$C$3:$C$6"))
    ws.print_area = "A1:H6"
    ws.page_setup.orientation = "landscape"
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 72
    for col in "CDEFGH":
        ws.column_dimensions[col].width = 16

    notes = wb.create_sheet("Notes")
    notes["A1"] = "Synthetic SECI fixture for cross-document / traceability tests."
    notes.merge_cells("A1:D1")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()

"""Phase 6: DataICD deterministic rules + applicability vs conformity separation."""

from __future__ import annotations

import uuid

from app.core.enums import ChecklistAnswerState
from app.services.checklist_catalog import CATALOG_DIR, load_catalog_file
from app.services.dataicd_rules import parse_icd_records, run_dataicd_rule
from app.services.evidence import EvidenceItem, EvidencePack
from app.services.excel_export import ExportRowWrite, fill_workbook_copy


def _pack(texts: list[str], *, doc_type: str = "data_icd") -> EvidencePack:
    pid = uuid.uuid4()
    vid = uuid.uuid4()
    items = []
    for i, t in enumerate(texts):
        items.append(
            EvidenceItem(
                evidence_id=f"ev_{i}",
                document_version_id=vid,
                project_id=pid,
                quote=t,
                locator={"kind": "unit", "doc_type": doc_type, "extra": {"doc_type": doc_type}},
                revision="1",
                content_sha256="sha",
            )
        )
    return EvidencePack(
        project_id=pid,
        document_version_id=vid,
        version_label="1",
        content_sha256="sha",
        items=items,
        items_by_version={str(vid): items},
        doc_types_by_version={str(vid): doc_type},
    )


GOOD_ROWS = [
    "SIG-001 | Airspeed | float | m/s",
    "SIG-002 | Altitude | int32 | ft",
    "PARAM-010 | Mode | enum | N/A",
]


def test_data_icd_catalog_scaffold():
    data = load_catalog_file(key="data_icd")
    assert data["key"] == "data_icd"
    assert data["template_gap"]["claimed_template_file_count"] == 0
    assert (CATALOG_DIR / "fixtures" / "data_icd_synthetic_v1.xlsx").is_file()
    cols = data["cell_mapping"]["columns"]
    assert cols["is_applicable"] == "C"
    assert cols["answer"] == "D"
    assert cols["is_applicable"] != cols["answer"]
    for item in data["items"]:
        assert "is_applicable" in item["excel_cells"]
        assert item["excel_cells"]["is_applicable"] != item["excel_cells"]["answer"]


def test_parse_and_unique_ids_yes():
    pack = _pack(GOOD_ROWS)
    records = parse_icd_records(pack)
    assert len(records) == 3
    result = run_dataicd_rule(
        "dataicd_unique_ids", {}, pack, applicability_yes=True
    )
    assert result.state == ChecklistAnswerState.YES


def test_duplicate_ids_no():
    pack = _pack(
        [
            "SIG-001 | Airspeed | float | m/s",
            "SIG-001 | Dup | float | m/s",
        ]
    )
    result = run_dataicd_rule(
        "dataicd_unique_ids", {}, pack, applicability_yes=True
    )
    assert result.state == ChecklistAnswerState.NO


def test_missing_records_insufficient_not_no():
    pack = _pack(["This document has no ICD table rows."])
    result = run_dataicd_rule(
        "dataicd_all_records_have_id", {}, pack, applicability_yes=True
    )
    assert result.state == ChecklistAnswerState.INSUFFICIENT_EVIDENCE


def test_all_have_datatype_and_units_yes():
    pack = _pack(GOOD_ROWS)
    r1 = run_dataicd_rule(
        "dataicd_all_records_have_datatype", {}, pack, applicability_yes=True
    )
    r2 = run_dataicd_rule(
        "dataicd_all_records_have_units", {}, pack, applicability_yes=True
    )
    assert r1.state == ChecklistAnswerState.YES
    assert r2.state == ChecklistAnswerState.YES


def test_applicability_gate_sets_na_not_written_as_conformity_into_is_applicable(tmp_path):
    pack = _pack(["unrelated text"], doc_type="source")
    app = run_dataicd_rule(
        "dataicd_applicability",
        {
            "applicable_doc_types": ["icd", "data_icd"],
            "icd_marker_patterns": ["(?i)\\bICD\\b"],
        },
        pack,
        applicability_yes=None,
        primary_doc_type="source",
    )
    assert app.state == ChecklistAnswerState.NO

    gated = run_dataicd_rule(
        "dataicd_unique_ids", {"gate_on_applicability": True}, pack, applicability_yes=False
    )
    assert gated.state == ChecklistAnswerState.NA

    # Export: Is Applicable=No, Answer=NA — not conformity Yes into Is Applicable
    fixture = CATALOG_DIR / "fixtures" / "data_icd_synthetic_v1.xlsx"
    dest = tmp_path / "out.xlsx"
    meta = fill_workbook_copy(
        fixture,
        dest,
        item_writes=[
            ExportRowWrite(
                item_key="DICD-Q01",
                cells={"C3": "No", "D3": "NA", "F3": "not applicable"},
                state_used="NA",
            ),
            ExportRowWrite(
                item_key="DICD-Q03",
                cells={"C5": "No", "D5": "NA"},
                state_used="NA",
            ),
        ],
        reference_rows=[],
        cell_mapping={"sheet": "Checklist"},
    )
    assert meta["cells_written"] >= 4
    from openpyxl import load_workbook

    wb = load_workbook(dest)
    ws = wb["Checklist"]
    assert ws["C3"].value == "No"  # applicability
    assert ws["D3"].value == "NA"  # conformity
    assert ws["C3"].value != "Yes"  # did not write conformity Yes into Is Applicable
    wb.close()


def test_applicability_yes_when_icd_markers_present():
    pack = _pack(
        ["Interface Control Document ICD listing follows", *GOOD_ROWS],
        doc_type="source",
    )
    result = run_dataicd_rule(
        "dataicd_applicability",
        {
            "applicable_doc_types": ["icd", "data_icd"],
            "icd_marker_patterns": ["(?i)\\bICD\\b", "(?i)interface control"],
        },
        pack,
        applicability_yes=None,
        primary_doc_type="source",
    )
    assert result.state == ChecklistAnswerState.YES

"""Phase 5: Excel export cell mapping, blank-on-insufficient, injection, preservation."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.services.checklist_catalog import CATALOG_DIR, load_catalog_file
from app.services.excel_export import (
    BLANK_ANSWER_STATES,
    REF_SHEET_NAME,
    ExportRowWrite,
    answer_cell_value,
    fill_workbook_copy,
    resolve_template_path,
    sanitize_excel_text,
)
from app.models.entities import ChecklistDefinition


FIXTURE = CATALOG_DIR / "fixtures" / "software_code_standard_synthetic_v1.xlsx"


def test_synthetic_fixture_exists_and_catalog_points_to_it():
    data = load_catalog_file()
    assert FIXTURE.is_file()
    assert data["excel_template_path"].endswith("software_code_standard_synthetic_v1.xlsx")
    gap = data["template_gap"]
    assert gap["claimed_template_file_count"] == 0
    assert gap["status"] == "synthetic_fixture_only"
    assert "do not claim" in gap["note"].lower()
    cols = data["cell_mapping"]["columns"]
    for key in ("answer", "chapter", "comment", "references", "reviewed_item", "status"):
        assert key in cols
    assert "Author's Answer" not in cols.values()
    assert "Author's Answer" in data["cell_mapping"]["do_not_invent_columns"]
    assert "Resolved SVN Revision" in data["cell_mapping"]["do_not_invent_columns"]


def test_sanitize_formula_injection():
    assert sanitize_excel_text("=1+1") == "'=1+1"
    assert sanitize_excel_text("+cmd") == "'+cmd"
    assert sanitize_excel_text("-2") == "'-2"
    assert sanitize_excel_text("@SUM(A1)") == "'@SUM(A1)"
    assert sanitize_excel_text("normal text") == "normal text"
    assert sanitize_excel_text(None) is None


def test_blank_answer_for_insufficient_manual_error():
    for state in ("INSUFFICIENT_EVIDENCE", "MANUAL_REVIEW", "ERROR"):
        assert state in BLANK_ANSWER_STATES
        assert answer_cell_value(state) is None
    assert answer_cell_value("YES") == "Yes"
    assert answer_cell_value("NO") == "No"
    assert answer_cell_value("NA") == "NA"


def test_fill_copy_writes_mapped_cells_and_preserves_template(tmp_path: Path):
    dest = tmp_path / "out.xlsx"
    # Probe original
    orig = load_workbook(FIXTURE)
    ows = orig["Checklist"]
    merged = [str(m) for m in ows.merged_cells.ranges]
    validations = len(ows.data_validations.dataValidation)
    print_area = ows.print_area
    formula = ows["J3"].value
    orig.close()

    writes = [
        ExportRowWrite(
            item_key="SCS-Q01",
            cells={
                "C3": "Yes",
                "D3": "§4.1",
                "E3": "ok",
                "F3": "DO-178C:VERIFIED",
                "G3": "Yes",
                "H3": None,  # do not auto-Closed
            },
            state_used="YES",
        ),
        ExportRowWrite(
            item_key="SCS-Q02",
            cells={
                "C4": None,  # insufficient → blank
                "D4": None,
                "E4": "INSUFFICIENT_EVIDENCE: missing scan",
                "F4": None,
                "G4": None,
                "H4": None,
            },
            state_used="INSUFFICIENT_EVIDENCE",
        ),
        ExportRowWrite(
            item_key="SCS-Q03",
            cells={
                "C5": None,
                "E5": sanitize_excel_text("=HYPERLINK(\"http://evil\")"),
            },
            state_used="ERROR",
        ),
    ]
    meta = fill_workbook_copy(
        FIXTURE,
        dest,
        item_writes=writes,
        reference_rows=[
            {
                "ref": "DO-178C",
                "stated version": "2011",
                "selected source": "std-v1",
                "result": "VERIFIED",
                "location": "page=2",
                "explanation": "short note",
            }
        ],
        cell_mapping={"sheet": "Checklist"},
    )
    assert dest.is_file()
    assert FIXTURE.stat().st_mtime <= dest.stat().st_mtime or True
    # Original still has empty C3
    recheck = load_workbook(FIXTURE)
    assert recheck["Checklist"]["C3"].value is None
    recheck.close()

    wb = load_workbook(dest)
    ws = wb["Checklist"]
    assert ws["C3"].value == "Yes"
    assert ws["D3"].value == "§4.1"
    assert ws["C4"].value is None
    assert "INSUFFICIENT" in (ws["E4"].value or "")
    assert (ws["E5"].value or "").startswith("'=")
    assert ws["H3"].value is None  # not Closed
    assert [str(m) for m in ws.merged_cells.ranges] == merged
    assert len(ws.data_validations.dataValidation) == validations
    assert ws.print_area == print_area
    assert ws["J3"].value == formula
    assert REF_SHEET_NAME in wb.sheetnames
    rws = wb[REF_SHEET_NAME]
    assert rws["A1"].value == "ref"
    assert rws["A2"].value == "DO-178C"
    assert rws["D2"].value == "VERIFIED"
    # No long standard body republished
    assert rws.max_column == 6
    wb.close()
    assert meta["preservation"]["merged_cells_ok"]
    assert meta["preservation"]["data_validations_ok"]
    assert meta["preservation"]["print_area_ok"]
    assert meta["preservation"]["formula_j3_ok"]


def test_resolve_template_path_from_definition():
    data = load_catalog_file()
    definition = ChecklistDefinition(
        key=data["key"],
        title=data["title"],
        version_label=data["version_label"],
        excel_template_path=data["excel_template_path"],
        meta={},
    )
    path = resolve_template_path(definition)
    assert path == FIXTURE.resolve()

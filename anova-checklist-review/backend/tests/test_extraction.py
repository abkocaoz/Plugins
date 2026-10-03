from __future__ import annotations

import io

from openpyxl import Workbook

from app.services.extraction import detect_format, extract_csv, extract_document, extract_xlsx


def test_detect_format():
    assert detect_format("a.pdf") == "pdf"
    assert detect_format("a.docx") == "docx"
    assert detect_format("a.xlsx") == "xlsx"
    assert detect_format("a.csv") == "csv"
    assert detect_format("main.c") == "code"
    assert detect_format("weird.bin") == "unknown"


def test_extract_csv_with_locators():
    data = "id,text\n1,hello\n2,world\n".encode("utf-8")
    result = extract_csv(data)
    assert result.status in {"ok", "partial"}
    assert result.units
    assert result.units[0].locator.kind == "csv_row"
    assert result.units[0].locator.row == 1


def test_extract_code_line_locators():
    data = "\n".join(f"line {i}" for i in range(1, 100)).encode("utf-8")
    result = extract_document(data, "sample.py")
    assert result.format == "code"
    assert result.units
    assert result.units[0].locator.line_start == 1
    assert result.units[0].locator.line_end is not None


def test_extract_xlsx_cells():
    wb = Workbook()
    ws = wb.active
    ws.title = "ICD"
    ws["A1"] = "Signal"
    ws["B1"] = "Rate"
    ws["A2"] = "CLK"
    ws["B2"] = "10 Hz"
    buf = io.BytesIO()
    wb.save(buf)
    result = extract_xlsx(buf.getvalue())
    assert result.status == "ok"
    assert any(u.locator.sheet == "ICD" for u in result.units)
    assert any(u.text == "10 Hz" for u in result.units)


def test_unsupported_records_issue():
    result = extract_document(b"\x00\x01\x02", "blob.bin")
    assert result.status == "failed"
    assert result.issues[0].code == "UNSUPPORTED_FORMAT"
